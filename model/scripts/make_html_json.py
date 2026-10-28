# 由建置列號對照（data/rowmap_v02.json）產生 carvana/html.json（共用產生器 tools/company_html 的設定檔）。
# 設定檔文字不得含阿拉伯數字（年度、日期除外）；所有數字以儲存格參照取自 xlsx。
import json, os
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RM = json.load(open(os.path.join(HERE, "data", "rowmap_v02.json"), encoding="utf-8"))
R, P, C = RM["R"], RM["P"], RM["C"]
IN = RM["ROWMAP"]["In"]; QR = RM["ROWMAP"]["Q"]; CRM = RM["ROWMAP"]["Credit"]; K = CRM["K"]
VR = RM["ROWMAP"]["Val"]["VR"]; VS = RM["ROWMAP"]["Val"]["VS"]; RC = RM["ROWMAP"]["Recon"]; SE = RM["ROWMAP"]["Sens"]
yc = lambda y: get_column_letter(3 + y - 2025)
Y0, Y1 = 2026, 2031
YC = [yc(y) for y in range(Y0, Y1 + 1)]


def rng(sheet, row):
    return f"@{sheet}!{yc(Y0)}{row}:{yc(Y1)}{row}"


def sc(sheet, key, src):
    return {s: rng(sheet, src[s][key]) for s in "ABC"}


def inrow(k, fmt, derived_label, derived_ref):
    r = IN[k]
    return [f"@In!B{r}", [f"@In!E{r}|{fmt}", " ", f"@In!C{r}"], f"@In!G{r}|tag", [derived_label, derived_ref]]


LAST = str(max(int(k) for k in RC))
import openpyxl as _ox
_wb = _ox.load_workbook(os.path.join(HERE, "dist", sorted(x for x in os.listdir(os.path.join(HERE, "dist")) if x.endswith(".xlsx"))[-1]), data_only=True)
RFMT = {}
for _k, _r in RC.items():
    _u = _wb["Recon"][f"B{_r}"].value
    RFMT[str(_r)] = {c: {"US$M": "n0", "輛": "n0", "US$": "n0", "%": "p1", "x": "n2"}.get(_u, "n2") for c in "CDE"}
B = R["B"]; PB = P["B"]; CB = C["B"]; KB = K["B"]; S1 = CRM["S1"]; S2 = CRM["S2"]; S3 = CRM["S3"]; S4 = CRM["S4"]; S5 = CRM["S5"]; S6 = CRM["S6"]
fmt_th = {}
for k in ("g", "gos_m"):
    fmt_th[str(B[k])] = {c: ("p1" if k == "g" else "p2") for c in YC}
cfg = {
    "company": "Carvana",
    "ticker": "CVNA",
    "thesis": "@README!A2",
    "question": "信用循環轉差時，Carvana 能否不靠新股與新債維持收支平衡？債務對 EBITDA 會不會越過評等門檻？",
    "scenarios": [{"key": "A", "name": "情境 A（樂觀）"}, {"key": "B", "name": "情境 B（基準）"}, {"key": "C", "name": "情境 C（信用壓力）"}],
    "sections": [
        {"id": "thesis", "title": "命題與核心問題", "blocks": [
            {"type": "kv", "items": [
                [f"@Val!A{VR['rec']}", f"@Val!C{VR['rec']}"],
                [f"@Val!A{VR['tp']}", f"@Val!C{VR['tp']}|n2"],
                [f"@Val!A{VR['px']}", f"@Val!C{VR['px']}|n2"],
                [f"@Val!A{VR['up']}", f"@Val!C{VR['up']}|sp1"],
                [f"@Val!A{VR['irr_w']}", f"@Val!C{VR['irr_w']}|p1"],
                [f"@Val!A{VR['wacc']}", f"@Val!C{VR['wacc']}|p1"],
                ["@CF!A4", ["@CF!C4|n0", "／", "@CF!D4|n0", "／", "@CF!E4|n0"]],
            ]},
            {"type": "h3", "text": "重點觀察"},
            {"type": "list", "items": ["@README!A5", "@README!A6", "@README!A7", "@README!A8", "@README!A9"]},
            {"type": "kv", "items": [["會計年度與口徑", "@README!A12"], ["營收分部層", "@README!A13"], ["與專案既有模型的關係", "@README!A14"]]},
        ]},
        {"id": "drivers", "title": "驅動因子 → 推導量（單位經濟）", "blocks": [
            {"type": "table", "head": ["輸入", "基準值", "標記", "推導量"], "rows": [
                inrow("q4_seq", "p1", "FY2026 零售台數（情境 B）：", f"@Thesis!{yc(2026)}{B['u']}|n0"),
                inrow("g27_B", "p1", "FY2027 零售台數（情境 B）：", f"@Thesis!{yc(2027)}{B['u']}|n0"),
                inrow("gpr27_B", "n0", "FY2027 零售毛利（情境 B，百萬美元）：", f"@Thesis!{yc(2027)}{B['gp_ret']}|n0"),
                inrow("gos27_B", "p2", "FY2027 貸款出售利得（情境 B，百萬美元）：", f"@Thesis!{yc(2027)}{B['gos']}|n0"),
                inrow("gos27_C", "p2", "FY2027 貸款出售利得（情境 C，百萬美元）：", f"@Thesis!{yc(2027)}{R['C']['gos']}|n0"),
                inrow("sga_var", "n0", "FY2027 SG&A 每台（情境 B）：", f"@Thesis!{yc(2027)}{B['sga_pu']}|n0"),
                inrow("dcnl_C", "p1", "情境 C 損失傳導（利得率）：", f"@Credit!C{S2['l']}|p2"),
                inrow("dspr_C", "p1", "情境 C 利差傳導（利得率）：", f"@Credit!C{S2['d']}|p2"),
                inrow("ally_cap", "n0", "已承諾去化 ÷ FY2027 發放（情境 B）：", f"@Credit!{yc(2027)}{KB['take']}|p0"),
            ]},
            {"type": "table", "caption": "命題模組：零售台數 × 每台單位經濟（情境 B）", "grid": True,
             "head": [f"@Thesis!A{B['hdr']}", f"@Thesis!B{B['hdr']}"] + [f"@Thesis!{c}{B['hdr']}" for c in [yc(2025)] + YC],
             "rowgen": {"sheet": "Thesis", "rows": f"{B['g']}-{B['gcx']}", "cols": ["A", "B", yc(2025)] + YC, "fmts": ["", ""] + ["n0"] * (len(YC) + 1),
                        "row_fmt": {**{str(B["g"]): {c: "p1" for c in [yc(2025)] + YC}}, **{str(B["gos_m"]): {c: "p2" for c in [yc(2025)] + YC}}}}},
            {"type": "note", "parts": ["@Thesis!A1"]},
            {"type": "note", "parts": ["@Thesis!A2"]},
        ]},
        {"id": "pl", "title": "三情境損益、收支平衡與信用風險", "blocks": [
            {"type": "pl", "years": f"@PL!{yc(Y0)}3:{yc(Y1)}3", "metrics": [
                {"name": "零售台數（輛）", "fmt": "n0", "rows": sc("Thesis", "u", R), "chart": True},
                {"name": "合併營收（百萬美元）", "fmt": "n0", "rows": sc("PL", "rev", P)},
                {"name": "總毛利每台（美元）", "fmt": "n0", "rows": sc("Thesis", "gpu", R)},
                {"name": "貸款出售利得（百萬美元）", "fmt": "n0", "rows": sc("PL", "sub_gos", P), "chart": True},
                {"name": "調整後 EBITDA（百萬美元）", "fmt": "n0", "rows": sc("PL", "ebitda", P), "chart": True},
                {"name": "調整後 EBITDA 率", "fmt": "p1", "rows": sc("PL", "ebitda_m", P)},
                {"name": "貸款出售利得 ÷ 調整後 EBITDA", "fmt": "p1", "rows": sc("PL", "gos_sh", P)},
                {"name": "淨利（百萬美元）", "fmt": "n0", "rows": sc("PL", "ni", P)},
                {"name": "每股盈餘（完全轉換股數，美元）", "fmt": "n2", "rows": sc("PL", "eps", P)},
            ]},
            {"type": "table", "caption": "營收線與損益（情境 B）", "grid": True,
             "head": ["@PL!A3", "@PL!B3"] + [f"@PL!{c}3" for c in [yc(2025)] + YC],
             "rowgen": {"sheet": "PL", "rows": f"{PB['rev_ret']}-{PB['gos_sh']}", "cols": ["A", "B", yc(2025)] + YC, "fmts": ["", ""] + ["n0"] * (len(YC) + 1),
                        "row_fmt": {str(PB[k]): {c: ("n2" if k == "eps" else "p1") for c in [yc(2025)] + YC} for k in ("rev_g", "oi_m", "ni_m", "ebitda_m", "gos_sh", "eps")}}},
            {"type": "table", "caption": "分部核對（三情境合計絕對差）", "grid": True, "head": ["@PL!A3", "@PL!B3"] + [f"@PL!{c}3" for c in [yc(2025)] + YC],
             "rowgen": {"sheet": "PL", "rows": "5-6", "cols": ["A", "B", yc(2025)] + YC, "fmts": ["", ""] + ["n1"] * (len(YC) + 1)}},
            {"type": "h3", "text": "收支模組（現金流、債務、收支關卡）"},
            {"type": "table", "grid": True, "head": ["@CF!A3", "@CF!C3", "@CF!D3", "@CF!E3"],
             "rowgen": {"sheet": "CF", "rows": "4-7", "cols": ["A", "C", "D", "E"], "fmts": ["", "n0", "n0", "n0"]}},
            {"type": "pl", "years": f"@PL!{yc(Y0)}3:{yc(Y1)}3", "metrics": [
                {"name": "自由現金流（百萬美元）", "fmt": "n0", "rows": sc("CF", "fcf", C), "chart": True},
                {"name": "收支關卡：自由現金流扣 TRA 與到期本金（百萬美元）", "fmt": "n0", "rows": sc("CF", "gate", C), "chart": True},
                {"name": "營運資金變動（百萬美元）", "fmt": "n0", "rows": sc("CF", "dwc", C)},
                {"name": "債務合計（百萬美元）", "fmt": "n0", "rows": sc("CF", "debt", C)},
                {"name": "利息費用淨額（百萬美元）", "fmt": "n0", "rows": sc("CF", "int", C)},
            ]},
            {"type": "table", "caption": "現金流與債務表（情境 C）", "collapse": True, "summary": "展開：情境 C 的完整現金流與債務表", "grid": True,
             "head": [f"@CF!A{C['C']['hdr']}", f"@CF!B{C['C']['hdr']}"] + [f"@CF!{c}{C['C']['hdr']}" for c in YC],
             "rowgen": {"sheet": "CF", "rows": f"{C['C']['inv']}-{C['C']['int']}", "cols": ["A", "B"] + YC, "fmts": ["", ""] + ["n0"] * len(YC)}},
            {"type": "note", "parts": ["@CF!A1"]},
            {"type": "h3", "text": "信用風險：公司債信指標"},
            {"type": "pl", "years": f"@PL!{yc(Y0)}3:{yc(Y1)}3", "metrics": [
                {"name": "債務對調整後 EBITDA（倍）", "fmt": "n2", "rows": sc("Credit", "lev", K), "chart": True},
                {"name": "債務對扣除出售利得的 EBITDA（倍）", "fmt": "n2", "rows": sc("Credit", "lev_xg", K), "chart": True},
                {"name": "債務加 TRA 對 EBITDA（倍）", "fmt": "n2", "rows": sc("Credit", "lev_tra", K)},
                {"name": "自由現金流對債務", "fmt": "p0", "rows": sc("Credit", "focf", K)},
            ]},
            {"type": "table", "caption": "評等門檻檢查（情境 C）", "grid": True, "head": ["@Credit!A3", "@Credit!B3"] + [f"@Credit!{c}3" for c in YC],
             "rowgen": {"sheet": "Credit", "rows": f"{K['C']['cov']},{K['C']['moody']},{K['C']['sp']}", "cols": ["A", "B"] + YC, "fmts": ["", ""] + ["n2"] * len(YC),
                        "row_fmt": {str(K["C"][k]): {c: "" for c in YC} for k in ("moody", "sp")}}},
            {"type": "kv", "items": [[f"@Credit!A{K['B']['take']}", [f"@Credit!{yc(2027)}{K['B']['take']}|p0", "（情境 B，FY2027）"]],
                                     [f"@Credit!A{K['B']['take2']}", [f"@Credit!{yc(2027)}{K['B']['take2']}|p0", "（情境 B，FY2027）"]],
                                     [f"@Credit!A{K['B']['cush']}", f"@Credit!{yc(2026)}{K['B']['cush']}|n1"],
                                     [f"@Q!A{QR['ttm_lev_xg']}", f"@Q!C{QR['ttm_lev_xg']}|n2"]]},
            {"type": "h3", "text": "信用風險：消費信用（管理組合）"},
            {"type": "table", "grid": True, "head": [f"@Credit!{c}{S1['hdr']}" for c in "ABCDEFGHI"],
             "rowgen": {"sheet": "Credit", "rows": f"{S1['bal']}-{S1['g']}", "cols": list("ABCDEFGHI"), "fmts": ["", ""] + ["n0"] * 7,
                        "row_fmt": {str(S1[k]): {c: "p2" for c in "CDEFGHI"} for k in ("ncor", "ncorp", "ncorn", "lag", "dq", "g")}}},
            {"type": "note", "parts": [f"@Credit!A{S1['hdr'] - 1}"]},
            {"type": "note", "parts": [f"@Credit!A{S1['g'] + 1}"]},
            {"type": "h3", "text": "信用風險：出售利得的傳導與資金斷點"},
            {"type": "table", "grid": True, "head": ["項目", "單位", "值"],
             "rowgen": {"sheet": "Credit", "rows": f"{S2['b']}-{S2['per1s']}", "cols": ["A", "B", "C"], "fmts": ["", "", "p2"],
                        "row_fmt": {str(S2["per1"]): {"C": "n0"}, str(S2["per1s"]): {"C": "p1"}}}},
            {"type": "table", "grid": True, "head": [f"@Credit!A{S3['h']}", f"@Credit!B{S3['h']}", f"@Credit!C{S3['h']}", f"@Credit!D{S3['h']}"],
             "rowgen": {"sheet": "Credit", "rows": f"{S3['o']}-{S3['whl']},{S3['on']}-{S3['runway']}", "cols": ["A", "B", "C", "D"], "fmts": ["", "", "n0", "n0"],
                        "row_fmt": {str(S3[k]): {"C": "n1", "D": "n1"} for k in ("m1", "runway")}}},
            {"type": "note", "parts": [f"@Credit!A{S3['runway'] + 1}"]},
            {"type": "h3", "text": "信用風險：盈餘品質、關係人與清償順位"},
            {"type": "table", "grid": True, "head": [f"@Credit!{c}{S4['hdr']}" for c in "ABCDEFGHIJK"],
             "rowgen": {"sheet": "Credit", "rows": f"{S4['hdr'] + 1}-{S4['lev_xg']}", "cols": list("ABCDEFGHIJK"), "fmts": ["", ""] + ["p1"] * 9,
                        "row_fmt": {str(S4["ebitda_xg"]): {c: "n0" for c in "CDEFGHIJK"}, str(S4["lev_xg"]): {"K": "n2"}, str(S4["gos_m"]): {c: "p2" for c in "CDEFGHIJK"}}}},
            {"type": "note", "parts": [f"@Credit!A{S4['lev_xg'] + 1}"]},
            {"type": "table", "grid": True, "head": [f"@Credit!{c}{S5['hdr']}" for c in "ABCDE"],
             "rowgen": {"sheet": "Credit", "rows": f"{S5['a_cash']}-{S5['a_tot']}", "cols": list("ABCDE"), "fmts": ["", "", "n0", "p0", "n0"]}},
            {"type": "table", "grid": True, "head": [f"@Credit!{c}{S5['c_abf'] - 1}" for c in "ABCDE"],
             "rowgen": {"sheet": "Credit", "rows": f"{S5['c_abf']}-{S5['c_tra']}", "cols": list("ABCDE"), "fmts": ["", "", "n0", "n0", "n2"]}},
            {"type": "table", "caption": f"@Credit!A{S6['hdr'] - 1}", "grid": True, "head": [f"@Credit!{c}{S6['hdr']}" for c in "ACDEF"],
             "rowgen": {"sheet": "Credit", "rows": f"{S6['hdr'] + 1}-{S6['end']}", "cols": list("ACDEF"), "fmts": ["p1", "n2", "n2", "n2", "n2"]}},
            {"type": "note", "parts": ["@Credit!A2"]},
        ]},
        {"id": "valuation", "title": "估值", "blocks": [
            {"type": "bars", "title": "目標價與現價", "fmt": "n1", "items": [
                {"label": f"@Val!A{VR['tp']}", "value": f"@Val!C{VR['tp']}", "series": 0},
                {"label": "情境 A 目標價", "value": f"@Val!C{VR['s_tp']}", "series": 1},
                {"label": "情境 B 目標價", "value": f"@Val!D{VR['s_tp']}", "series": 1},
                {"label": "情境 C 目標價", "value": f"@Val!E{VR['s_tp']}", "series": 2},
                {"label": f"@Val!A{VR['an']}", "value": f"@Val!C{VR['an']}", "series": 1}],
             "ref": {"label": f"@Val!A{VR['px']}", "value": f"@Val!C{VR['px']}"}},
            {"type": "table", "grid": True, "head": [f"@Val!{c}{VR['sum_hdr']}" for c in "ACDE"],
             "rowgen": {"sheet": "Val", "rows": f"{VR['s_dcf']}-{VR['prob']}", "cols": ["A", "C", "D", "E"], "fmts": ["", "n2", "n2", "n2"],
                        "row_fmt": {str(VR["s_tvsh"]): {c: "p0" for c in "CDE"}, str(VR["s_tvm"]): {c: "n1" for c in "CDE"}, str(VR["s_irr"]): {c: "p1" for c in "CDE"}, str(VR["prob"]): {c: "p0" for c in "CDE"}}}},
            {"type": "kv", "items": [[f"@Val!A{VR[k]}", f"@Val!C{VR[k]}|{f}"] for k, f in (("ke", "p2"), ("wacc", "p2"), ("peer", "n2"), ("mult", "n2"), ("mkt_mult", "n1"), ("irr_w", "p1"))]},
            {"type": "table", "grid": True, "head": [f"@Val!{c}{VR['res_hdr']}" for c in "ACE"],
             "rowgen": {"sheet": "Val", "rows": f"{VR['tp']}-{VR['rec']}", "cols": ["A", "C", "E"], "fmts": ["", "n2", ""],
                        "row_fmt": {str(VR["up"]): {"C": "sp1"}, str(VR["req30"]): {"C": "n0"}, str(VR["b30"]): {"C": "n0"}, str(VR["req_u"]): {"C": "n0"}, str(VR["rec"]): {"C": ""}}}},
            {"type": "note", "parts": ["@Val!A1"]},
        ]},
        {"id": "recon", "title": "對帳：公司指引、共識、評等機構、放空報告", "blocks": [
            {"type": "table", "head": [f"@Recon!{c}3" for c in "ABCDEFG"],
             "rowgen": {"sheet": "Recon", "rows": f"{RC['0']}-{RC[LAST]}", "cols": list("ABCDEFG"), "fmts": ["", "", "n2", "n2", "n2", "", "tag"],
                        "row_fmt": RFMT}},
            {"type": "note", "parts": [f"@Recon!A{RC[LAST] + 2}"]},
            {"type": "note", "parts": ["@Recon!A1"]},
        ]},
        {"id": "sens", "title": "敏感度", "blocks": [
            {"type": "tornado", "title": "加權目標價相對基準：輸入取低值", "sheet": "Sens", "rows": f"{SE['first']}-{SE['last']}", "label_col": "A", "value_col": "F", "fmt": "sp1"},
            {"type": "tornado", "title": "加權目標價相對基準：輸入取高值", "sheet": "Sens", "rows": f"{SE['first']}-{SE['last']}", "label_col": "A", "value_col": "G", "fmt": "sp1"},
            {"type": "table", "collapse": True, "summary": "展開：全部案例（含情境 B 債務對 EBITDA、情境 C 收支關卡）", "grid": True,
             "head": [f"@Sens!{c}3" for c in "ACDFGHIJK"],
             "rowgen": {"sheet": "Sens", "rows": f"{SE['first']}-{SE['last']}", "cols": list("ACDFGHIJK"), "fmts": ["", "g", "g", "sp1", "sp1", "n2", "n2", "n0", "n0"]}},
            {"type": "note", "parts": ["@Sens!A1"]},
        ]},
        {"id": "status", "title": "資料狀態", "blocks": [
            {"type": "kv", "items": [["模型版本", ["$version", "（", "$date", "）"]], ["來源 xlsx", "$xlsx"], ["來源 xlsx 的雜湊指紋", "$sha"]]},
            {"type": "h3", "text": "已知限制"},
            {"type": "note", "parts": ["@README!A17"]},
            {"type": "h3", "text": "「未找到」清單"},
            {"type": "scan", "needle": "未找到", "collapse": True, "summary": "展開：標為「未找到」的輸入與來源", "sheets": {"In": ["B"], "Recon": ["A"]}, "empty": "沒有標為未找到的項目。"},
        ]},
    ],
}
json.dump(cfg, open(os.path.join(HERE, "html.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print("wrote html.json")
