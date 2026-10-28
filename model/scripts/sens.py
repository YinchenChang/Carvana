# 敏感度：一次動一個輸入（低／高），以 SENS_OVR 完整重建並用 LibreOffice 重算，讀回三個輸出。
# 用法：python3 carvana/scripts/sens.py --jobs 6   → 寫 carvana/data/sens_v03.json（再建置一次時寫入 Sens 頁）
import argparse, json, os, subprocess, sys, tempfile, shutil
from concurrent.futures import ThreadPoolExecutor
import openpyxl

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BUILD = os.path.join(HERE, "build_carvana.py")
OUT = os.path.join(HERE, "data", "sens_v03.json")
KEYS = ["q4_seq", "g27_B", "g28_B", "g30_B", "g35_B", "gpr27_B", "gprLT_B", "gos27_B", "gosLT_B", "sga_fix", "sga_var", "eff_B", "capcost", "util",
        "dcnl_C", "dspr_C", "dur", "etr_30", "ctr", "tra_yrs", "beta", "g_tv", "prem", "inv_d", "refi_sp", "gprLT_C", "gosLT_C", "asp_g"]


def run(ovr, tag, tmp):
    d = os.path.join(tmp, tag); os.makedirs(d, exist_ok=True)
    env = dict(os.environ, SENS_OVR=json.dumps(ovr), CVNA_ROWMAP=os.path.join(d, "rowmap.json"))
    subprocess.run([sys.executable, BUILD, os.path.join(d, "raw.xlsx")], env=env, check=True, capture_output=True)
    subprocess.run(["soffice", f"-env:UserInstallation=file://{d}/lo", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", os.path.join(d, "out"), os.path.join(d, "raw.xlsx")],
                   check=True, capture_output=True, timeout=300)
    rm = json.load(open(os.path.join(d, "rowmap.json")))
    wb = openpyxl.load_workbook(os.path.join(d, "out", "raw.xlsx"), data_only=True)
    tp = wb["Val"]["C%d" % rm["ROWMAP"]["Val"]["VR"]["tp"]].value
    lev = wb["Credit"]["E%d" % rm["ROWMAP"]["Credit"]["K"]["B"]["lev"]].value
    gate = wb["CF"]["E4"].value
    return {"tp": tp, "lev": lev, "gateC": gate}


def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--jobs", type=int, default=4); a = ap.parse_args()
    tmp = tempfile.mkdtemp(prefix="cvna_sens_")
    # 讀低／高值（In 頁 D、F 欄）
    base_dir = os.path.join(tmp, "base0"); os.makedirs(base_dir)
    subprocess.run([sys.executable, BUILD, os.path.join(base_dir, "raw.xlsx")], check=True, capture_output=True, env=dict(os.environ, SENS_OVR="{}"))
    ws = openpyxl.load_workbook(os.path.join(base_dir, "raw.xlsx"))["In"]
    lohi = {}
    for r in range(4, ws.max_row + 1):
        k = ws.cell(r, 1).value
        if k in KEYS:
            lohi[k] = (ws.cell(r, 4).value, ws.cell(r, 6).value)
    jobs = [("base", {})]
    for k in KEYS:
        lo, hi = lohi[k]
        jobs += [(f"{k}__lo", {k: lo}), (f"{k}__hi", {k: hi})]
    with ThreadPoolExecutor(a.jobs) as ex:
        res = dict(zip([j[0] for j in jobs], ex.map(lambda j: run(j[1], j[0], tmp), jobs)))
    out = {"base": res["base"], "cases": {k: {"lo_val": lohi[k][0], "hi_val": lohi[k][1], "lo": res[f"{k}__lo"], "hi": res[f"{k}__hi"]} for k in KEYS}}
    json.dump(out, open(OUT, "w"), ensure_ascii=False, indent=1)
    shutil.rmtree(tmp, ignore_errors=True)
    print("wrote", OUT, "base", out["base"])


if __name__ == "__main__":
    main()
